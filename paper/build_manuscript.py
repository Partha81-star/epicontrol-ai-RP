from pathlib import Path
import json,csv,hashlib,copy,math,html,re
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'D:\research_ papers')
WORK=ROOT/'parth-adi-yug-manya(epicontrol-ai)'/'epicontrol_work'
RES=WORK/'Epicontrol_AI'/'results'
FIG=WORK/'figures';FIG.mkdir(exist_ok=True)
OUT=ROOT/'parth-adi-yug-manya(epicontrol-ai)'/'EpiControl_AI_IEEE_Manuscript.docx'
summary=json.loads((RES/'summary.json').read_text())
extra=json.loads((RES/'additional_summary.json').read_text())
fontpath=Path(r'C:\Windows\Fonts\times.ttf')
def font(n):return ImageFont.truetype(str(fontpath),n)
def csvrows(name):
    with (RES/name).open() as f:return list(csv.DictReader(f))
COLORS=['#155b8d','#a84324','#28794c','#744c96','#5b626a']
def chart(name,series,ylabel,xlabel='Time (days)',xmax=180,ymax=None,xmin=0):
    im=Image.new('RGB',(1500,1050),'white');dr=ImageDraw.Draw(im)
    left,top,right,bottom=170,90,1440,790
    ymax=ymax or max(y for _,pts,_ in series for _,y in pts)*1.08
    ymax=max(ymax,.01)
    for j in range(5):
        yy=bottom-(bottom-top)*j/4
        dr.line((left,yy,right,yy),fill='#d9d9d9',width=2)
        value=ymax*j/4
        lab=f'{value:,.0f}' if ymax>10 else f'{value:.2f}'
        dr.text((10,yy-18),lab,font=font(36),fill='black')
        xx=left+(right-left)*j/4
        tick=xmin+(xmax-xmin)*j/4
        label=f'{tick:.3f}' if xmax<=1 else f'{tick:g}'
        dr.text((xx-25,bottom+15),label,font=font(36),fill='black')
    dr.line((left,top,left,bottom,right,bottom),fill='black',width=3)
    dr.text((left,15),ylabel,font=font(42),fill='black')
    dr.text((550,855),xlabel,font=font(42),fill='black')
    for idx,(label,pts,color) in enumerate(series):
        xy=[(left+(right-left)*(x-xmin)/(xmax-xmin),bottom-(bottom-top)*y/ymax) for x,y in pts]
        dr.line(xy,fill=color,width=6)
        lx=50+(idx%3)*480;ly=950+(idx//3)*50
        dr.line((lx,ly,lx+65,ly),fill=color,width=6)
        dr.text((lx+80,ly-20),label,font=font(34),fill='black')
    im.save(FIG/name,dpi=(450,450))

series=[]
for name,label,color in [('None','No restriction',COLORS[0]),('Threshold','Threshold',COLORS[1]),('Q11','Q seed 11',COLORS[2]),('Strict','Strict uniform',COLORS[3])]:
    rows=csvrows(f'trajectory_{name}.csv');series.append((label,[(float(r['day']),float(r['I'])) for r in rows],color))
chart('infection_curves.png',series,'Infectious population (people)')
series=[]
for s,color in zip((11,22,33),COLORS):
    rows=csvrows(f'training_seed{s}.csv');vals=[float(r['return']) for r in rows]
    pts=[(i+1,sum(vals[max(0,i-99):i+1])/len(vals[max(0,i-99):i+1])) for i in range(len(vals))]
    series.append((f'Seed {s}',pts,color))
# Training return is negative; plot magnitude so the generic graph keeps a zero baseline.
series=[(name,[(x,-y) for x,y in pts],color) for name,pts,color in series]
chart('training_curves.png',series,'Mean negative return (100 episodes)','Training episode',1200)
series=[]
rows=csvrows('sensitivity.csv')
for compliance,color in zip((.5,.7,.85,1.),COLORS):
    selected=[r for r in rows if r['policy']=='Q11' and float(r['compliance'])==compliance]
    series.append((f'Compliance {compliance:g}',[(float(r['beta']),float(r['peak_infectious'])) for r in selected],color))
chart('sensitivity.png',series,'Peak infectious population (people)','Transmission rate (per day)',.42,xmin=.24)

ref=ROOT/'Full-Paper-template.docx'
sha=hashlib.sha256(ref.read_bytes()).hexdigest()
(WORK/'artifact.md').write_text(f'''# IEEE manuscript layout contract
Reference: {ref}
SHA256: {sha}
Reference render: template-render/template.pdf and page-1.png through page-4.png; all four pages inspected.
Reference sections: 3; US Letter 8.5 by 11 inches; top and bottom 1 inch; left and right 1210 twips; columns 2 with gap 461 twips.
Styles retained: Normal Times New Roman 10 pt; Body Text Indent first line 0.17 inch, justified, single; Heading 1 10 pt small caps; Heading 2 10 pt italic; Caption 10 pt small caps; Author Data 12 pt; paper title 24 pt.
Slot contract: replace all sample title, authors, body, table, figure, biography, and reference content. Remove sample comments, author image, conference-specific footer and example links. Preserve style definitions, numbering and theme. The user requests IEEE format; add Roman section numbering and lettered subsections explicitly.
Intentional geometry changes: use two sections rather than the example's three, retaining a full-width title followed by continuous two-column body. Remove the third biography section. Title uses Word Title style based on paper title. Research figures and compact result tables replace the sample illustration and table. Black headings; no page numbers.
Final gates: source hash unchanged, 10 pt body, two equal columns, equations editable in Word, no unsupported old result figures, references numbered in order of first citation, every final page visually inspected in Microsoft Word PDF export and bundled Poppler PNGs. Canonical renderer attempted but unavailable because LibreOffice is not bundled on Windows.
''',encoding='utf-8')

d=Document(ref)
body=d._element.body
last=copy.deepcopy(body.sectPr)
for child in list(body):body.remove(child)
body.append(last)
for s in d.sections:
    for hdr in (s.header,s.footer,s.first_page_header,s.first_page_footer):
        for p in hdr.paragraphs:p.clear()
    s.different_first_page_header_footer=False
    cols=s._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'1')
    for e in list(s._sectPr):
        if e.tag in (qn('w:headerReference'),qn('w:footerReference')):s._sectPr.remove(e)
for name in ('Normal','Body Text Indent','Heading 1','Heading 2','Caption','Author Data','Abstract','paper title'):
    st=d.styles[name];st.font.color.rgb=RGBColor(0,0,0)
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
from docx.enum.style import WD_STYLE_TYPE
title=d.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH) if 'Title' not in d.styles else d.styles['Title']
title.base_style=d.styles['paper title'];title.font.name='Times New Roman';title.font.size=Pt(24)
title.font.color.rgb=RGBColor(0,0,0);title.paragraph_format.space_after=Pt(8)
title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ('Heading 1','Heading 2'):
    st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(10);st.paragraph_format.keep_with_next=True
    # Preserve appearance while removing template automatic list numbering.
    num=st.element.find('.//'+qn('w:numPr'))
    if num is not None:num.getparent().remove(num)

def para(text,style='Body Text Indent'):
    p=d.add_paragraph(text,style);p.paragraph_format.widow_control=True
    if style=='Body Text Indent':
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.first_line_indent=Inches(.17)
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    return p
def heading(text,level=1):
    p=para(text,f'Heading {level}');p.paragraph_format.keep_with_next=True
    num=OxmlElement('w:numPr');ilvl=OxmlElement('w:ilvl');ilvl.set(qn('w:val'),'0');nid=OxmlElement('w:numId');nid.set(qn('w:val'),'0');num.extend([ilvl,nid]);p._p.get_or_add_pPr().append(num)
    return p
def eq(text,number):
    p=para('', 'Normal');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
    om=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);om.append(r);p._p.append(om)
    rr=p.add_run(f'   ({number})');rr.font.size=Pt(9)
def fig(file,caption):
    p=para('','Normal');p.paragraph_format.space_before=Pt(6);p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(FIG/file),width=Inches(3.22))
    p=para(caption,'Caption');p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=False
    for r in p.runs:r.font.size=Pt(9);r.font.small_caps=False
def table(caption,headers,rows,widths):
    p=para(caption,'Caption');p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
    for r in p.runs:r.font.size=Pt(9);r.font.small_caps=False
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','bottom','left','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);marg=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'45');e.set(qn('w:type'),'dxa');marg.append(e)
            pr.append(marg)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAEAEA');pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                if ri<len(t.rows)-1:p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8);r.bold=ri==0
    para('','Normal').paragraph_format.space_after=Pt(3)

para('EpiControl AI\nA Reproducible Two Cohort SEIR Framework\nfor Adaptive Intervention Analysis','Title')
authors=[
    ('Prof Bhagyashri Thorat','Department of Electronics and\nTelecommunication','International Institute Of Information\nTechnology','Pune, India','email'),
    ('Parth Omprakash Bhad','Department of Computer Engineering','International Institute Of Information\nTechnology','Pune, India','parthbhad2@gmail.com'),
    ('Aaditya Shrikrishna Hingmire','Department of Information Technology','International Institute Of Information\nTechnology','Pune, India','aadityahingmire@gmail.com'),
    ('Yugraj Prabhakar Mangate','Department of Computer Engineering','International Institute Of Information\nTechnology','Pune, India','yugrajmngate@gmail.com'),
    ('Manasvi Manoj Yeole','Department of Computer Engineering','International Institute Of Information\nTechnology','Pune, India','manasviyeole@gmail.com'),
]
author_table=d.add_table(rows=2,cols=3)
author_table.alignment=WD_TABLE_ALIGNMENT.CENTER
author_table.autofit=False
for row in author_table.rows:
    for cell in row.cells:
        cell.width=Inches(2.38);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        tcpr=cell._tc.get_or_add_tcPr()
        borders=OxmlElement('w:tcBorders')
        for edge in ('top','bottom','left','right','insideH','insideV'):
            elem=OxmlElement(f'w:{edge}');elem.set(qn('w:val'),'nil');borders.append(elem)
        tcpr.append(borders)
        margins=OxmlElement('w:tcMar')
        for margin_name in ('top','bottom','left','right'):
            mar=OxmlElement(f'w:{margin_name}');mar.set(qn('w:w'),'55');mar.set(qn('w:type'),'dxa')
            margins.append(mar)
        tcpr.append(margins)
slots=[(0,0),(0,1),(0,2),(1,0),(1,1)]
for author,(ridx,cidx) in zip(authors,slots):
    name,dept,inst,place,email=author
    cell=author_table.cell(ridx,cidx)
    cell.text=''
    p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    pieces=[(name,False),(dept,False),(inst,False),(place,True),(email,False)]
    for idx,(text,is_italic) in enumerate(pieces):
        run=p.add_run(text)
        run.font.name='Times New Roman';run.font.size=Pt(9.5);run.italic=is_italic
        if idx<len(pieces)-1:run.add_break()
for p in author_table.cell(1,2).paragraphs:
    p.paragraph_format.space_after=Pt(0)
para('','Author Data').paragraph_format.space_after=Pt(3)
s=d.add_section(WD_SECTION_START.CONTINUOUS);cols=s._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'461')
abstract=('Abstract - Epidemic intervention models are useful only when their assumptions, policy costs, and evaluation boundaries are visible. '
'This paper presents EpiControl AI, a reproducible decision-support prototype that couples a two-cohort susceptible-exposed-infectious-recovered model with weekly tabular reinforcement learning and an interactive web interface. '
'The simulator represents 49,000 people as aggregate child and adult compartments. Three independently trained controllers are evaluated on 40 held-out synthetic scenarios over 180 days and compared with four fixed policies and a prevalence-triggered policy. '
'The pooled learned controllers produce a mean peak infectious population of 968 people, compared with 11,900 without restrictions, and reduce the mean restriction-cost proxy by 47.4% relative to continuous strict restrictions. '
'However, strict restrictions achieve a lower infection peak, and the learned controllers show substantial variation across training seeds. '
'A follow-up experiment that removes restrictions after day 180 produces renewed outbreaks, demonstrating that short-horizon control does not imply lasting elimination. '
'These findings support transparent scenario comparison under stated assumptions. The contact matrix, hospital-demand proxy, and policy-cost function require empirical calibration before the framework can inform operational decisions.')
p=para(abstract,'Abstract');p.paragraph_format.first_line_indent=Inches(0);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after=Pt(6)
para('Index Terms - Adaptive intervention, epidemic simulation, reinforcement learning, SEIR model.','Abstract')
heading('I. Introduction')
para('Epidemic response requires repeated decisions about contact reduction, timing, and the burden placed on different population groups. A restriction that reduces transmission may also interrupt work and education. The resulting trade-off changes as infections accumulate and compliance varies. Compartmental models provide an interpretable starting point for describing these changes [1], while age-dependent contact patterns motivate explicit population stratification [2].')
para('Reinforcement learning offers a way to select interventions from repeated interactions with a simulator. Earlier studies have already combined epidemic models and learned policies [3]-[6]. The practical research question addressed here is narrower: can a small, inspectable controller improve a stated synthetic objective, and what limitations become visible when the same experiment is examined across training seeds and beyond its original horizon?')
para('EpiControl AI addresses this question through a shared simulation engine for training, evaluation, and the user interface. Its contributions are a reproducible two-cohort environment, a paired comparison against fixed and threshold policies, and an explicit post-release assessment. The experimental population is represented by eight compartment values rather than individual agents. The study evaluates a computational prototype under synthetic assumptions; it does not estimate the effectiveness of any policy in a particular country.')
heading('II. Related Work')
para('The compartmental tradition established by Kermack and McKendrick [1] describes epidemic progression through population-level transitions. An exposed compartment introduces a latent period before infectiousness. In the present formulation, exposed individuals do not transmit; modelling presymptomatic infectiousness would require a separate transmitting compartment or an additional term in the force of infection.')
para('Prem et al. [2] projected age-specific contact matrices using survey and demographic information. Their work shows how population structure can inform mixing assumptions, while also identifying uncertainty when direct contact measurements are absent. EpiControl AI uses a deliberately illustrative two-cohort matrix. Its values are not taken from those country estimates, and stronger adult mixing is an experimental assumption rather than an observed epidemiological finding.')
para('Kompella et al. [3] developed an agent-based simulator with interactions at specific locations and studied mitigation policies that balance economic activity and hospital capacity. Their modelling resolution supports interventions that alter individual behaviour and local contacts. The present aggregate simulator offers a smaller implementation that is easier to inspect, with substantially less behavioural detail.')
para('Ohi et al. [4] examined reinforcement learning for balancing epidemic outcomes and economic effects. Rao [5] combined a SEIRD environment with a deep double recurrent Q-network and alternative reward definitions. Jain et al. [6] considered a SIR framework with restriction and vaccination decisions, emphasizing health and economic trade-offs. Together, these studies establish that coupling epidemic simulation with reinforcement learning is an existing research direction.')
para('The controller in this study uses tabular Q-learning [7]. The contribution lies in documenting an executable comparison, exposing seed variability, and testing what happens after restrictions are removed. Published results from other simulators are not treated as directly comparable numerical baselines because population definitions, rewards, horizons, and calibration differ.')
heading('A. Contact Structure and Cohort Heterogeneity',2)
para('Diary-based measurements by Mossong et al. [8] provide direct evidence that contact patterns vary with age and setting. Davies et al. [9] modelled age-dependent susceptibility and clinical manifestations of COVID-19. These papers motivate cohort structure, but neither supplies the two-by-two contact coefficients used in this experiment.')
para('Recent analyses have examined which contact entries most influence model outputs [10], who acquires infection from whom during early spread in Belgium [11], age-specific contact contributions in Germany [12], and age-dependent transmission under suppression measures in Shanghai [13]. Together, they show that a fixed two-cohort matrix is an interpretable starting assumption rather than a transferable estimate of contact behaviour.')
heading('B. Intervention Evidence and Behaviour',2)
para('Studies of early pandemic control estimated effects across European countries [14], ranked combinations of measures across territories [15], analysed large-scale anti-contagion policies [16], and inferred effects of particular government interventions [17]. These studies use different designs and outcomes. Their estimates cannot be inserted directly as the effect size of the synthetic actions in Table I.')
para('Contact rates also respond to interventions, as studied in a protocol for updating contact matrices [18]. Reviews of model-informed measures [19] and UK interventions [20] place estimates within a wider, heterogeneous evidence base. An analysis of self-protection in Germany [21] shows why behaviour cannot be represented solely by formal policy intensity. A European modeller survey [22] identifies data and decision-process needs, while a multi-country study jointly examines vaccination and non-pharmaceutical interventions [23]. Post-pandemic Italian contact surveys [24] and a stochastic study of rural US epidemics [25] further illustrate differences in settings and scale.')
heading('C. Adaptive and Learning Based Control',2)
para('Extensions beyond a single homogeneous decision environment include a spatial model with variant-specific transmission [26], connected community and correctional-facility networks [27], real-time policy responses [28], and regional influenza interventions [29]. Hierarchical reinforcement learning has been investigated for multiple control modes [30], and semi-connected SEIQR regions have been used to study outbreak management [31]. These models represent forms of heterogeneity that the present tabular controller omits.')
para('Simulation with deep reinforcement learning [32] and large-scale agent-based pandemic control [33] explore richer action or population representations. Their algorithms and evaluation settings differ from this study. The present controller is deliberately small enough that its state bins, reward, training seeds, and failure under parameter shift can be examined directly.')
heading('D. Hospital Demand and Capacity',2)
para('Healthcare planning studies have combined epidemic forecasts with patient allocation [34], optimized bed schedules across facilities [35], and generated local bed-demand projections on request [36]. A pragmatic bed-occupancy forecast was evaluated with health-system data [37], and near-term case and hospitalisation forecasts have been built for Aotearoa New Zealand [38]. These works model operational quantities that cannot be recovered from infectious prevalence alone.')
para('Forecasting ICU demand using a two-stage approach [39] and hospital admissions through a time-series model [40] provide additional alternatives. Vaccination and healthcare demand have been analysed together [41], while a capacity-based intervention approach links projected demand to control [42]. EpiControl AI therefore labels 0.05 times current infectious prevalence as a demand proxy and does not report it as observed occupancy.')
heading('E. Calibration and Uncertainty',2)
para('Data-based model fitting and uncertainty analysis are essential for a claim about real-world decisions. Approximate Bayesian computation has been accelerated for agent-based calibration [43], multiscale methods have been proposed for early-outbreak assessment [44], and neural calibration has been combined with uncertainty estimates [45]. A review of epidemiological and molecular surveillance [46] discusses the information available to support monitoring.')
para('The US Scenario Modeling Hub has been evaluated as a way to inform response under uncertainty [47]. Ensemble sub-epidemic forecasts [48], Bayesian estimation of epidemiological parameters [49], and multiscale spatial modelling under uncertain data [50] offer distinct routes to uncertainty treatment. This literature explains why the parameter grid reported in Section VI is a stress test, not a calibrated confidence interval for a specific population.')
heading('III. Model and Controller')
heading('A. Cohort Dynamics',2)
para('Let g denote the child or adult cohort. Each cohort contains susceptible S_g, exposed E_g, infectious I_g, and recovered R_g people, with constant size N_g. Births, deaths, reinfection, migration, and vaccination are omitted. Recovered individuals remain immune for the simulated horizon. The force of infection depends on the infectious share in both cohorts:')
eq('λ_g = β f_g Σ_h C_gh f_h I_h / N_h',1)
eq('dS_g/dt = −λ_g S_g',2)
eq('dE_g/dt = λ_g S_g − σ E_g',3)
eq('dI_g/dt = σ E_g − γ I_g',4)
eq('dR_g/dt = γ I_g',5)
para('Here β is a baseline transmission coefficient, σ is the latent-to-infectious transition rate, and γ is the recovery rate, all expressed per day. C_gh is a dimensionless relative mixing coefficient. A policy assigns restriction intensity u_g to each cohort, and f_g = 1 − c u_g, where c is compliance. Applying the reduction on both sides of an interaction represents a symmetric contact-reduction assumption. It is not an empirically estimated intervention effect.')
para('The child and adult fractions are 0.3 and 0.7. The contact rows are (0.8, 0.35) and (0.15, 1.2), respectively. Cross-cohort contacts satisfy reciprocity because N_child C_child,adult = N_adult C_adult,child. Initial exposed and infectious counts are allocated in proportion to cohort size. The ordinary differential equations are integrated with fourth-order Runge-Kutta using two half-day steps per day; outputs are recorded daily.')
heading('B. Weekly Actions and Observation',2)
para('Decisions are made every seven days, with a shorter final interval when necessary. Table I defines the available actions. The restriction-cost proxy is C(a) = Σ_g (N_g/N)u_g². It describes relative intervention intensity; it is not a measured financial loss, a percentage of economic activity, or a validated welfare index.')
table('TABLE I. Synthetic intervention actions',['Action','Child u','Adult u','Cost'],[
 ['None','0.00','0.00','0.0000'],['Moderate','0.35','0.35','0.1225'],['Targeted','0.15','0.65','0.3025'],['Strict','0.85','0.85','0.7225']],[1.04,.69,.69,.8])
para('The controller observes a discretized state containing time, susceptible share, exposed share, infectious share, and the previous action. Time is divided into six 30-day bins. Susceptible-share thresholds are 0.5 and 0.9; exposed and infectious thresholds are 0.002 and 0.02. The resulting table has 648 state bins and four actions, or 2,592 action values. This compression loses cohort detail and continuous parameter information, so it is an approximate state representation and does not establish a fully observed Markov decision process.')
heading('C. Reward and Q Update',2)
para('The hospital-demand proxy is H_d = 0.05 I_d, where I_d is total infectious population on day d. Synthetic capacity K is 100. Hospital admissions, clinical delay, and length of stay are not separately simulated. The weekly reward penalizes new infections, daily excess demand, intervention cost, and changes in restriction intensity:')
eq('B_t = Σ_d [Δ_d + max(H_d−K,0)] / N',6)
eq('r_t = −B_t − 0.002 n_t C(a_t) − 0.0002 V_t',7)
para('The daily sum in (6) covers the current decision interval, whose length is n_t days; Δ_d is the reduction in susceptible population that day. V_t is the mean absolute change in the two restriction intensities from the preceding action. All terms are dimensionless. Their relative weights encode a chosen synthetic objective rather than a consensus policy preference.')
eq('y_t = r_t + δ max_b Q(s′,b)',8)
eq('Q(s,a) ← (1−η) Q(s,a) + η y_t',9)
para('The learning rate η is 0.1 and the discount factor δ is 0.95 per weekly decision. Terminal transitions use reward alone as the target. Exploration follows an epsilon-greedy rule: epsilon starts at 1.0, is multiplied by 0.997 after each episode, and has a floor of 0.05. Q values start at zero. Greedy evaluation freezes the table and resolves ties by choosing the lowest action index. These implementation choices are recorded because they can affect the learned policy [7].')
heading('IV. Software Architecture')
para('The prototype follows a monitoring, modelling, and managing workflow. Monitoring displays simulated compartment trajectories and derived indicators. Modelling uses a shared SEIR environment for every policy. Managing exposes fixed, threshold, and learned policies through a Flask interface. The code derives from the EpiControl AI repository [51] and includes a repaired model module, explicit validation, and repeatable experiment scripts.')
para('The web application provides scenario controls for population, transmission, compliance, initial conditions, capacity, and horizon. A scenario endpoint returns numerical metrics and daily trajectories, and a comparison endpoint evaluates policies under identical inputs. Cohort curves make the mixing assumptions visible. No external surveillance service is connected, and geographic infection rates are not inferred from population weights. Model output is presented as synthetic scenario data throughout the interface.')
para('Experiment outputs include training returns, saved Q-tables, parameter manifests, evaluation records, sensitivity records, and daily example trajectories. Three training seeds are retained independently. The published-source baseline is repository commit 8d5cf6a964f2c008a163cbbbac40d8fda04cfaad; the experiment uses the accompanying local revision rather than that commit alone. The source and data package must therefore accompany the manuscript for reproduction.')
heading('V. Experimental Design')
heading('A. Training and Held Out Scenarios',2)
para('Each controller is trained for 1,200 episodes, using seeds 11, 22, and 33. At the start of each episode, β is sampled uniformly from 0.24 to 0.36 per day and compliance from 0.7 to 1.0. Initial exposed counts are sampled uniformly from the integers 10 through 50, and initial infectious counts from 5 through 20. Population, contact matrix, latent rate, recovery rate, capacity, and reward weights remain fixed. Dynamics within an episode are deterministic; randomness comes from scenario generation and action exploration.')
table('TABLE II. Shared experiment assumptions',['Parameter','Value'],[
 ['Population','49,000'],['Child and adult shares','0.30 and 0.70'],['Training and evaluation horizon','180 days'],['Latent transition σ','0.20 per day'],['Recovery γ','0.10 per day'],['Hospital fraction and capacity','0.05 and 100'],['RK4 step and policy interval','0.5 day and 7 days'],['Episodes per training seed','1,200'],['Held-out scenario seeds','10,000 through 10,039']],[1.85,1.37])
para('Evaluation uses 40 independently generated scenarios from the same parameter ranges, with seeds 10,000 through 10,039. The same scenario is applied to each of the three frozen Q-tables and each baseline. The fixed policies continuously apply none, moderate uniform, adult-targeted, or strict uniform restrictions. The threshold policy applies the targeted action when total infectious prevalence exceeds 0.002, and otherwise applies no restrictions; its decision interval is also seven days.')
heading('B. Metrics and Additional Checks',2)
para('Reported metrics are peak infectious population, attack rate, days above the hospital-demand proxy capacity, mean restriction cost, and undiscounted cumulative reward. Attack rate is (N − S_T)/N and therefore includes people initially exposed or infectious. Peak values include day zero. Learned-policy summaries pool 120 evaluations, representing three training seeds applied to the same 40 scenarios; these records are correlated rather than 120 independent scenarios.')
para('A paired comparison first averages learned-policy return across the three training seeds for each scenario and then subtracts the threshold return. A scenario bootstrap resamples the 40 paired differences 10,000 times using seed 20,261,007. This interval describes scenario variability conditional on the three trained controllers; it does not fully characterize uncertainty over possible training runs.')
para('Sensitivity analysis evaluates the seed-11 controller at β values 0.24, 0.30, 0.36, and 0.42, and compliance values 0.50, 0.70, 0.85, and 1.00, using 20 initially exposed and 10 initially infectious people. Values 0.42 and 0.50 extend beyond the training ranges. Integration checks compare steps of 0.5, 0.25, and 0.125 day. A separate extension applies each learned policy or strict restrictions until day 180, removes all restrictions, and continues to day 365 under unchanged disease parameters.')
heading('VI. Results')
heading('A. Policy Comparison',2)
para('Table III reports mean outcomes. No restrictions produce a mean peak of 11,899.6 infectious people and a 96.91% attack rate. The pooled learned policies reduce the mean peak to 968.0, a 91.9% reduction relative to the no-restriction mean. Their mean attack rate is 12.42%. Strict uniform restrictions yield a much lower peak of 24.5 and an attack rate of 0.147%, with a restriction cost of 0.7225 throughout the horizon.')
rows=[]
for name,label in [('None','None'),('Moderate','Moderate'),('Targeted','Targeted'),('Strict','Strict'),('Threshold','Threshold'),('Q pooled','Q pooled')]:
    ss=summary[name];rows.append([label,f"{ss['peak_infectious']['mean']:,.1f}",f"{ss['attack_rate']['mean']*100:.2f}",f"{ss['mean_restriction_cost']['mean']:.3f}",f"{ss['return']['mean']:.3f}"])
table('TABLE III. Mean outcomes over the evaluation scenarios',['Policy','Peak','Attack %','Cost','Return'],rows,[.75,.66,.68,.54,.59])
para('The learned policies have a mean restriction cost of 0.3797, which is 47.4% lower than strict uniform restrictions. Their mean return is −0.2638, compared with −0.2608 for strict restrictions and −0.5253 for the threshold rule. Thus the pooled learned policy does not outperform strict restrictions on mean combined return. Its paired return improvement over the threshold rule is 0.2614, with a scenario-bootstrap 95% interval of [0.2391, 0.2845]. This comparison is specific to the reward weights and selected baseline.')
para('The mean number of days above the demand-proxy capacity is 2.03 for the pooled learned policies, 26.50 for the threshold policy, and zero for strict restrictions. The pooled peak standard deviation is 983.5 people, and attack-rate standard deviation is 10.87 percentage points. These large dispersions indicate that a mean curve alone would conceal important variation. Figure 1 shows the first held-out scenario as an example rather than an average or a selected best case.')
fig('infection_curves.png','Fig. 1. Daily infectious population for held-out scenario seed 10000. The Q curve uses training seed 11; all curves share the same parameters.')
heading('B. Training Variability',2)
para('Table IV separates the three training seeds. Their mean peak infectious populations range from 465.6 to 1,528.6, and their mean returns range from −0.2065 to −0.3372. The seed-11 controller exceeds demand-proxy capacity for a mean of 4.83 days, whereas the seed-33 controller does not exceed it in these scenarios. Reporting only the strongest seed would overstate reliability.')
table('TABLE IV. Learned-policy variation across training seeds',['Seed','Mean peak','Mean attack %','Mean return'],[
 [s,f"{summary[f'Q{s}']['peak_infectious']['mean']:.1f}",f"{summary[f'Q{s}']['attack_rate']['mean']*100:.2f}",f"{summary[f'Q{s}']['return']['mean']:.4f}"] for s in (11,22,33)],[.4,.88,1.02,.92])
para('Figure 2 displays trailing 100-episode means of negative training return; lower values represent better outcomes. The curves describe exploratory training on changing scenarios, rather than greedy-policy evaluation on a fixed validation set. They do not establish convergence to an optimal controller.')
fig('training_curves.png','Fig. 2. Trailing mean negative training return for three seeds. Early windows contain all available episodes up to a maximum of 100.')
heading('C. Sensitivity and Numerical Checks',2)
para('The learned controller is sensitive to transmission and compliance (Fig. 3). At β = 0.42 and compliance 0.50, its peak reaches 12,877.1 infectious people and its attack rate reaches 99.06%. At the same transmission coefficient and compliance 1.00, the peak is 2,701.8 and the attack rate is 35.00%. These stress tests use one controller and fixed initial conditions; they indicate limited robustness rather than population-level estimates.')
fig('sensitivity.png','Fig. 3. Seed-11 controller evaluated across transmission coefficients and compliance values. Points are deterministic scenario outputs joined for readability.')
para('Across the main evaluation records, maximum population-conservation error is approximately 1.2 × 10⁻¹⁰ people. For the default no-restriction scenario, the infectious peak is 12,372.0951 using a 0.5-day integration step and 12,372.0938 using a 0.125-day step. Relative peak difference is approximately 1.1 × 10⁻⁷, with the peak recorded on day 65 in both cases. Numerical agreement supports implementation consistency, not empirical validity.')
heading('D. Outcomes After Restriction Removal',2)
para('Removing restrictions after day 180 produces renewed spread in every controller group. The mean post-release peaks are 8,303.3, 9,657.4, and 10,615.6 for training seeds 11, 22, and 33. The corresponding day-365 attack rates are 95.08%, 96.01%, and 96.42%. Strict restrictions also defer rather than permanently remove risk: their mean post-release peak is 11,490.9 and day-365 attack rate is 89.66%.')
para('The extension assumes deterministic residual infection, unchanged transmissibility, persistent immunity, and no vaccination or additional response. It excludes stochastic extinction, imported cases, and changing clinical care. Its purpose is to show that a controller rewarded over 180 days can postpone infections beyond the evaluation boundary. The extension does not identify a suitable real-world reopening policy.')
heading('VII. Discussion and Limitations')
para('The results support two conclusions within the stated environment. First, the learned controllers improve the synthetic objective relative to a simple prevalence-triggered targeted policy while using less restriction intensity than continuous strict control. Second, their performance varies across training seeds, deteriorates under some parameter shifts, and does not imply durable epidemic elimination. These findings make the evaluation boundary as important as the apparent reduction in peak infections.')
para('Several assumptions restrict interpretation. Contact coefficients and cohort fractions are illustrative, and policy intensity reduces both sides of a contact. The observed cohort differences therefore depend on the chosen mixing structure. The hospital proxy has no admission or discharge dynamics, and economic cost is represented only by squared restriction intensity. Mortality, vaccination, clinical severity, behavioural adaptation, reporting delay, and unequal intervention burdens are not modelled.')
para('The controller uses compressed aggregate observations and only three training seeds. The benchmark set includes a basic threshold rule but no optimized threshold, model-predictive controller, or continuous-action reinforcement learning algorithm. The reward weights are fixed, and preference sensitivity would require retraining. A policy can exploit the finite horizon by delaying infections; the post-release results show why longer-horizon and terminal-risk objectives deserve separate evaluation.')
para('Before operational use, the model would need calibration against traceable case and admission data, independent temporal validation, a clinically grounded hospital compartment, uncertainty in observation and transmission, and explicit treatment of fairness and feasibility. Interface testing with intended users would also be needed. The current application is suited to educational scenario exploration and reproducible computational experiments.')
heading('VIII. Conclusion')
para('EpiControl AI provides an executable two-cohort SEIR simulator, a weekly tabular Q-learning controller, and a common interface for policy comparison. In held-out synthetic scenarios, learned policies substantially reduce infection peaks relative to no restrictions and reduce the restriction-cost proxy relative to strict uniform control. Their mean combined return is close to that of strict restrictions, with considerable variability across training seeds. Sensitivity and post-release experiments identify failures that would be missed by a single short demonstration. The framework provides a transparent basis for further modelling and validation, with claims limited to its recorded synthetic experiments.')
heading('Acknowledgment')
para('OpenAI Codex, accessed October 7, 2026, assisted with drafting and restructuring Sections I-VIII, checking bibliographic records, and generating the revised software, experiment scripts, and figures. Figures 1-3 were rendered from recorded simulation outputs rather than generated as illustrative result images. This assistance does not constitute authorship; responsibility for scientific verification, attribution, and approval of the submitted article remains with the human authors.')
heading('References')
refs=[
'[1] W. O. Kermack and A. G. McKendrick, “A contribution to the mathematical theory of epidemics,” Proc. R. Soc. Lond. A, vol. 115, no. 772, pp. 700-721, 1927, doi: 10.1098/rspa.1927.0118.',
'[2] K. Prem, A. R. Cook, and M. Jit, “Projecting social contact matrices in 152 countries using contact surveys and demographic data,” PLoS Comput. Biol., vol. 13, no. 9, Art. no. e1005697, 2017, doi: 10.1371/journal.pcbi.1005697.',
'[3] V. Kompella et al., “Reinforcement learning for optimization of COVID-19 mitigation policies,” arXiv:2010.10560, 2020, doi: 10.48550/arXiv.2010.10560.',
'[4] A. Q. Ohi, M. F. Mridha, M. M. Monowar, and M. A. Hamid, “Exploring optimal control of epidemic spread using reinforcement learning,” Sci. Rep., vol. 10, Art. no. 22106, 2020, doi: 10.1038/s41598-020-79147-8.',
'[5] I. Rao, “Modeling and optimization of epidemiological control policies through reinforcement learning,” arXiv:2402.06640, 2024, doi: 10.48550/arXiv.2402.06640.',
'[6] M. Jain, Z. Uddin, and W. Ibrahim, “SIR-RL: Reinforcement learning for optimized policy control during epidemiological outbreaks in emerging market and developing economies,” arXiv:2404.08423, 2024, doi: 10.48550/arXiv.2404.08423.',
'[7] C. J. C. H. Watkins and P. Dayan, “Q-learning,” Mach. Learn., vol. 8, pp. 279-292, 1992, doi: 10.1007/BF00992698.'
]
def addref(num,authors,title,journal,year,doi,volume=None,issue=None,article=None):
    location=f', vol. {volume}' if volume else ''
    location+=f', no. {issue}' if issue else ''
    location+=f', Art. no. {article}' if article else ''
    refs.append(f'[{num}] {authors}, “{title},” {journal}{location}, {year}, doi: {doi}.')
addref(8,'J. Mossong et al.','Social contacts and mixing patterns relevant to the spread of infectious diseases','PLoS Med.',2008,'10.1371/journal.pmed.0050074',5,3,'e74')
addref(9,'N. G. Davies et al.','Age-dependent effects in the transmission and control of COVID-19 epidemics','Nat. Med.',2020,'10.1038/s41591-020-0962-9',26)
selected=json.loads((WORK/'literature_selected.json').read_text(encoding='utf8'))
def ieee_authors(author_string):
    names=[s.strip() for s in author_string.split(',') if s.strip()]
    def one(raw):
        bits=raw.split();initials=bits[-1]
        return ' '.join(c+'.' for c in initials if c.isalpha())+' '+' '.join(bits[:-1])
    if len(names)>3:return one(names[0])+' et al.'
    if len(names)==1:return one(names[0])
    if len(names)==2:return one(names[0])+' and '+one(names[1])
    return one(names[0])+', '+one(names[1])+', and '+one(names[2])
def selected_refs(group,start):
    items=[r for r in selected if r['group']==group]
    for i,r in enumerate(items,start):
        info=r['journalInfo'];journal=info.get('journal',{}).get('isoabbreviation') or info.get('journal',{}).get('title') or r.get('journalTitle')
        title=html.unescape(re.sub('<[^>]+>','',r['title'])).rstrip('.')
        addref(i,ieee_authors(r['authorString']),title,journal,r['pubYear'],r['doi'],info.get('volume'),info.get('issue'))
selected_refs('contact_age',10)
addref(14,'S. Flaxman et al.','Estimating the effects of non-pharmaceutical interventions on COVID-19 in Europe','Nature',2020,'10.1038/s41586-020-2405-7',584)
addref(15,'N. Haug et al.','Ranking the effectiveness of worldwide COVID-19 government interventions','Nat. Hum. Behav.',2020,'10.1038/s41562-020-01009-0',4)
addref(16,'S. Hsiang et al.','The effect of large-scale anti-contagion policies on the COVID-19 pandemic','Nature',2020,'10.1038/s41586-020-2404-8',584)
addref(17,'J. M. Brauner et al.','Inferring the effectiveness of government interventions against COVID-19','Science',2021,'10.1126/science.abd9338',371,6531)
selected_refs('intervention',18)
selected_refs('epidemic_rl',26)
selected_refs('hospital',34)
selected_refs('calibration',43)
refs.append('[51] Partha81-star, “Epicontrol_AI,” GitHub repository, baseline commit 8d5cf6a964f2c008a163cbbbac40d8fda04cfaad. Accessed: Oct. 7, 2026. [Online]. Available: https://github.com/Partha81-star/Epicontrol_AI.')
assert len(refs)==51 and all(s.startswith(f'[{i}]') for i,s in enumerate(refs,1))
for txt in refs:
    p=para(txt,'Normal');p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.first_line_indent=Inches(-.15);p.paragraph_format.left_indent=Inches(.15)
    for r in p.runs:r.font.size=Pt(9)
# Word balances the final two columns when a continuous final section follows.
end=d.add_section(WD_SECTION_START.CONTINUOUS)
end._sectPr.find(qn('w:cols')).set(qn('w:num'),'2')
settings=d.settings.element
hyphen=OxmlElement('w:autoHyphenation');hyphen.set(qn('w:val'),'true');settings.append(hyphen)
d.core_properties.title='EpiControl AI A Reproducible Two Cohort SEIR Framework for Adaptive Intervention Analysis'
d.core_properties.subject='Synthetic epidemic intervention experiment'
d.core_properties.author='EpiControl AI research team'
d.core_properties.comments=''
d.save(OUT)
assert hashlib.sha256(ref.read_bytes()).hexdigest()==sha
print(OUT)
print('Paragraph words',sum(len(p.text.split()) for p in d.paragraphs))
